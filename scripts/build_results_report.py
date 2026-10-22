from pathlib import Path
import csv
import statistics as stats
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OTHER = ROOT.parent / 'Resnet50test'
OUT = ROOT / 'reports'

def read(path):
    with open(path, encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))

u = read(ROOT/'results/cross_validation/unet_segmentation_cv_results.csv')
up = read(ROOT/'results/cross_validation/fusion_pipeline_cv_results.csv')
d = read(OTHER/'results/deeplabv3_pipeline_kfold_cv_results.csv')
ds = read(OTHER/'results/deeplabv3_kfold_cv_results.csv')
comparison = read(OTHER/'results/unet_vs_deeplabv3_comparison_v2.csv')
preds = read(ROOT/'results/fusion_pipeline/full_pipeline_test_results.csv')
classes = ['developing','flowering','fruiting','seeding']
cm = [[sum(r['true_stage']==a and r['pred_stage']==b for r in preds) for b in classes] for a in classes]
pr=[]; re=[]; f1=[]
for i in range(4):
    p=cm[i][i]/sum(row[i] for row in cm); r=cm[i][i]/sum(cm[i])
    pr.append(p); re.append(r); f1.append(2*p*r/(p+r))
acc=sum(cm[i][i] for i in range(4))/len(preds)

def summary(values, scale=1, digits=4):
    v=[float(x)*scale for x in values]
    return f'{stats.mean(v):.{digits}f} ± {stats.stdev(v):.{digits}f}'

def col(rows,key,scale=1,digits=4):
    return summary([r[key] for r in rows],scale,digits)

parts=[]
def p(text): parts.append(text.strip()+'\n')
def table(headers,rows):
    p('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+'\n'.join('| '+' | '.join(map(str,r))+' |' for r in rows))

p('# Tomato Leaf Segmentation and Growth-Stage Classification\n\nConsolidated Results Report\n\nPrepared: 30 September 2026')
p('This document separates segmentation-model performance, complete sequential pipeline performance, and cross-validation results. It consolidates saved results from the U-Net and Resnet50test projects. No models were retrained or inference rerun for this report. Fold summaries and U-Net fixed-test classification metrics were recalculated from saved CSV files. DeepLabV3 fixed-split macro metrics were transcribed from saved notebook output. Source identifiers [S1–S10] are resolved in Section 8.')
p('## 1. Study scope and experimental organization')
p('The segmentation task labels pixels as leaf/plant foreground or background. The classification task assigns an image to developing, flowering, fruiting, or seeding. The sequential pipeline predicts a mask, applies it to the RGB image, and sends the masked image to a ResNet50 classifier. DeepLabV3–ResNet50 names the segmentation network and its backbone; the downstream ResNet50 classifier is a separate model. These results concern sequential masking, not learned concatenation of LPF and classifier features.')
table(['Experiment','Evaluation design','What is measured'],[
['Fixed-split segmentation','Saved trained segmenter; existing test split','Pixel overlap and pixel classification'],
['Fixed-split pipeline','Saved segmenter and classifier; 160 test images','Four-stage image classification'],
['U-Net five-fold study','822 development images; five U-Nets; fixed 160-image test set','Segmenter variation and downstream variation with a fixed classifier'],
['DeepLabV3 segmenter-only CV','979 annotated images; five outer test folds','Generalization of the segmentation training procedure'],
['DeepLabV3 full-pipeline CV','979 annotated images; both models retrained per outer fold','Generalization of the sequential two-model training procedure']])
p('The original U-Net split contains 687 training, 135 validation, and 160 test images (982 total). DeepLabV3 CV records contain 979 usable annotated images after excluding samples without usable annotations. The exact three-image difference and cross-project identity should be reconciled before describing the CV studies as matched experiments. [S1, S6, S7]')
p('### 1.1 Metric conventions')
p('IoU and Dice are reported on the 0–1 scale. Classification accuracy, macro precision, macro recall, and macro F1 are percentages. Macro classification metrics average the four stages equally. Segmentation macro metrics average background and foreground equally. Foreground IoU is TP/(TP+FP+FN); binary Dice is 2TP/(2TP+FP+FN). Soft Dice uses probabilities and is a different measurement. Binary Dice equals foreground F1 only when masks and aggregation are identical.')
p('All fold summaries use an unweighted arithmetic mean and sample standard deviation (denominator K−1). A percentage metric’s SD is in percentage points. SD describes fold/run spread, not a confidence interval. DeepLabV3 CSV values were saved to four decimal places; derived summaries inherit this precision limit. Mean IoU cannot generally be converted into mean Dice after averaging.')
p('## 2. Segmentation-model results: fixed-split evaluation')
p('### 2.1 Saved comparison using binary masks')
p('Table 1 reports the saved cross-model comparison artifact. Foreground IoU and foreground Dice use the same class definition within this table. Preprocessing variants remain separate. These are single-checkpoint results and have no fold-level SD. [S2]')
table(['Segmentation model / preprocessing','Foreground IoU','Foreground Dice','Both-class mean IoU'],[[r['segmenter'],f"{float(r['seg_foreground_iou']):.4f}",f"{float(r['seg_foreground_dice']):.4f}",f"{float(r['seg_both_class_mean_iou']):.4f}"] for r in comparison])
p('In this saved comparison, U-Net’s foreground overlap was higher than DeepLabV3’s. This is a descriptive checkpoint comparison, not evidence of a statistically significant or cross-validated advantage. Both-class mean IoU includes background and is not interchangeable with foreground IoU.')
p('### 2.2 Original U-Net standalone evaluation')
table(['Split','Images','Binary IoU','Soft Dice'],[['Validation',135,'0.7241','0.8127'],['Test',160,'0.6650','0.7572']])
p('These values come from the original export metadata and use direct image resizing. Dice here is probability-based soft Dice; it must not be substituted for the binary Dice in Table 1. The difference between test soft Dice 0.7572 and binary Dice 0.7844 is therefore not a transcription error. [S1, S10]')
table(['Test growth stage','Images','Binary IoU','Soft Dice'],[['Developing',41,'0.6868','0.7736'],['Flowering',60,'0.6313','0.7289'],['Fruiting',21,'0.5186','0.6250'],['Seeding',38,'0.7756','0.8575']])
p('Fruiting had the lowest overlap in this evaluation; seeding had the highest. The saved results identify the performance pattern but do not establish its cause.')
p('### 2.3 Additional DeepLabV3 segmentation evaluations')
p('The fixed-split training notebook records test both-class mean IoU of 0.7548. A later saved fixed-split metrics cell records pixel accuracy 0.9010, macro precision 0.8466, macro recall 0.8626, and macro F1 0.8542; foreground F1/binary Dice is 0.7715. These are separate saved evaluations from the comparison artifact’s foreground Dice of 0.7565 and both-class IoU of 0.7507. They are retained here for provenance, not assembled into a synthetic result row. [S7, cells 17, 24 and 34; S2]')
p('## 3. Full sequential pipeline results: fixed-split evaluation')
p('### 3.1 Main fixed-split classification results')
p('Each row below describes its own saved evaluation of the segment → mask → classify sequence. The U-Net row is recalculated from 160 per-image predictions. The DeepLabV3 row uses the later saved standalone recomputation and its matching macro-metric output. [S3, S7 cell 34, S8 cell 12]')
table(['Pipeline','Accuracy (%)','Macro precision (%)','Macro recall (%)','Macro F1 (%)'],[
['U-Net → ResNet50',f'{acc*100:.2f}',f'{stats.mean(pr)*100:.2f}',f'{stats.mean(re)*100:.2f}',f'{stats.mean(f1)*100:.2f}'],
['DeepLabV3 → ResNet50 (recomputed)','87.50','85.46','88.57','86.12']])
p('U-Net correctly classified 143 of 160 images; DeepLabV3’s recomputed accuracy corresponds to 140 of 160. These evaluations reuse separately trained checkpoints. They do not establish a controlled architecture comparison or statistical superiority.')
p('### 3.2 Stage-wise fixed-split classification')
table(['Pipeline','Stage','Precision (%)','Recall (%)','F1 (%)','Test images'],
[['U-Net',c.title(),f'{pr[i]*100:.2f}',f'{re[i]*100:.2f}',f'{f1[i]*100:.2f}',sum(cm[i])] for i,c in enumerate(classes)] +
[['DeepLabV3','Developing','95.35','100.00','97.62',41],['DeepLabV3','Flowering','91.67','73.33','81.48',60],['DeepLabV3','Fruiting','54.84','80.95','65.38',21],['DeepLabV3','Seeding','100.00','100.00','100.00',38]])
p('Both fixed-split evaluations classified all 38 seeding images correctly. Fruiting had the lowest class F1. This describes the observed test set, not guaranteed deployment performance.')
p('### 3.3 U-Net fixed-split confusion matrix')
table(['True / predicted']+[c.title() for c in classes],[[c.title()]+cm[i] for i,c in enumerate(classes)])
p('Rows are true labels and columns are predictions. Thirteen of the 17 U-Net pipeline errors were exchanges between flowering and fruiting; four were exchanges between developing and flowering. [S3]')
p('### 3.4 Earlier preprocessing comparison: separate result series')
table(['Saved comparison variant','Pipeline accuracy (%)'],[[r['segmenter'],f"{float(r['pipeline_accuracy'])*100:.2f}"] for r in comparison])
p('This comparison records DeepLabV3 accuracy 84.38%, native-preprocessing U-Net accuracy 78.75%, and crop-aligned U-Net accuracy 89.38%. Do not attach the later DeepLabV3 macro metrics for the 87.50% run to the earlier 84.38% run. The crop-aligned and native U-Net results show that pipeline preprocessing matters alongside segmentation overlap. The comparison CSV stores rounded accuracies; the 89.38% value in Section 3.1 is independently calculated from exact predictions. [S2, S3]')
p('## 4. Cross-validation: segmentation models')
p('### 4.1 U-Net training and evaluation protocol')
p('Five growth-stage-stratified folds were formed from the 822 training-plus-validation images (seed 42). Each fresh U-Net trained for 30 epochs. Folds 1–2 used 657 training and 165 validation images; folds 3–5 used 658 and 164. The best validation-soft-Dice checkpoint was selected per fold. Each checkpoint was then evaluated on the fixed 160-image test set. The baseline used 60 epochs and a learning-rate scheduler, whereas the CV loop used 30 epochs without that scheduler. [S4, S9]')
p('The following foreground IoU values are derived per fold from saved pooled leaf F1 as IoU = F1/(2−F1), then summarized. They were not directly saved as IoU measurements. Binary Dice is the saved leaf F1 because both refer to the same pooled thresholded pixels. These values differ in aggregation from the per-image fixed-split metrics in Section 2.')
uiou=[float(r['test_leaf_f1'])/(2-float(r['test_leaf_f1'])) for r in u]
table(['Fold','Test foreground IoU (derived)','Test foreground Dice/F1','Test pixel accuracy (%)'],[[i+1,f'{uiou[i]:.4f}',f"{float(r['test_leaf_f1']):.4f}",f"{float(r['test_accuracy'])*100:.2f}"] for i,r in enumerate(u)]+[['Mean ± SD',summary(uiou),col(u,'test_leaf_f1'),col(u,'test_accuracy',100,2)]])
table(['U-Net segmentation metric','Fold-validation subsets: mean ± SD','Fixed test: mean ± SD'],[
['Foreground precision',col(u,'val_leaf_precision'),col(u,'test_leaf_precision')],
['Foreground recall',col(u,'val_leaf_recall'),col(u,'test_leaf_recall')],
['Foreground Dice/F1',col(u,'val_leaf_f1'),col(u,'test_leaf_f1')],
['Pixel accuracy',col(u,'val_accuracy'),col(u,'test_accuracy')]])
p('Fold-validation subsets also selected the checkpoint, so their scores are model-selection measurements. Fixed-test results measure variation among five trained segmenters on the same images; they are not five disjoint outer test folds.')
p('### 4.2 DeepLabV3 segmenter-only five-fold CV')
p('This separate experiment used ordinary, non-stratified KFold on 979 annotated images, with 10 epochs per fold and a 15% inner validation split. Folds 1–4 used 665/118/196 training/validation/test images; fold 5 used 666/118/195. Test IoU averages background and plant IoU within a batch and then averages batches equally. Pixel precision, recall, and F1 are macro averages over the two classes. [S5, S7 Section 9]')
table(['Fold','Both-class IoU','Pixel accuracy (%)','Macro precision (%)','Macro recall (%)','Macro F1 (%)'],[[r['fold'],f"{float(r['test_iou']):.4f}"]+[f'{float(r[k])*100:.2f}' for k in ['test_accuracy','test_precision_macro','test_recall_macro','test_f1_macro']] for r in ds]+[['Mean ± SD',col(ds,'test_iou')]+[col(ds,k,100,2) for k in ['test_accuracy','test_precision_macro','test_recall_macro','test_f1_macro']]])
p('Foreground-only Dice was not saved in this CV CSV. The recorded macro F1 is the mean of background and foreground F1, not foreground Dice. It must not be relabeled as leaf Dice or derived from the averaged IoU.')
p('### 4.3 DeepLabV3 segmentation within full-pipeline CV')
p('The full-pipeline study uses a different set of stratified fold runs. Both the segmenter and classifier were retrained; segmentation results from those runs are reported separately below. [S6]')
table(['Metric','Mean ± SD'],[['Both-class, batch-averaged IoU',col(d,'seg_test_iou')],['Pixel accuracy (%)',col(d,'seg_test_accuracy',100,2)],['Macro pixel precision (%)',col(d,'seg_test_precision_macro',100,2)],['Macro pixel recall (%)',col(d,'seg_test_recall_macro',100,2)],['Macro pixel F1 (%)',col(d,'seg_test_f1_macro',100,2)]])
p('The 0.7466 segmenter-only CV IoU and 0.7489 full-pipeline CV segmentation IoU are both recorded results from different runs. Neither is directly comparable with U-Net’s derived pooled foreground IoU without aligning the class definition, aggregation, data splits, and training protocol.')
p('## 5. Cross-validation: full sequential pipeline results')
p('### 5.1 Protocol differences')
p('U-Net: each of the five fold segmenters was paired with the same unchanged masked-image ResNet50 and evaluated on the same 160 test images. This is downstream robustness under segmenter retraining, not full two-model cross-validation. [S9]')
p('DeepLabV3: five stratified outer folds were formed from all 979 usable annotated images, including images from the original test partition. Each fold reserved 15% of its outer training partition for inner validation. A fresh pretrained segmenter trained for 10 epochs and a fresh pretrained classifier trained for 12 epochs. The classifier trained and validated on ground-truth-masked images; the outer test pipeline used predicted masks. Both models were selected using inner validation. This is full sequential training-procedure CV, although the networks were trained separately rather than jointly. [S6, S7 Sections 10]')
p('### 5.2 Pipeline summary table')
table(['Pipeline','Accuracy (%)','Macro precision (%)','Macro recall (%)','Macro F1 (%)'],[
['U-Net → fixed ResNet50']+[col(up,k,100,2) for k in ['accuracy','macro_precision','macro_recall','macro_f1']],
['DeepLabV3 → ResNet50; both retrained']+[col(d,k,100,2) for k in ['pipeline_test_accuracy','pipeline_test_precision_macro','pipeline_test_recall_macro','pipeline_test_f1_macro']]])
p('Table note: values are mean ± sample SD across five runs/folds. Evaluation protocols differ; this table is a descriptive summary and does not establish a matched CV ranking. Segmentation IoU is intentionally reported separately in Section 4.')
p('### 5.3 U-Net pipeline results by fold')
table(['Fold','Accuracy (%)','Macro precision (%)','Macro recall (%)','Macro F1 (%)'],[[i+1]+[f'{float(r[k])*100:.2f}' for k in ['accuracy','macro_precision','macro_recall','macro_f1']] for i,r in enumerate(up)])
p('U-Net pipeline accuracy ranged from 83.13% to 89.38%. The five evaluations contain 160 unique test images, not 800 independent test cases. The mean accuracy is lower than the fixed-split baseline, but the different training budgets and training subsets prevent attributing the difference solely to cross-validation.')
p('### 5.4 DeepLabV3 full-pipeline results by fold')
table(['Fold','Test images','Accuracy (%)','Macro precision (%)','Macro recall (%)','Macro F1 (%)'],[[r['fold'],r['n_test']]+[f'{float(r[k])*100:.2f}' for k in ['pipeline_test_accuracy','pipeline_test_precision_macro','pipeline_test_recall_macro','pipeline_test_f1_macro']] for r in d])
p('DeepLabV3 full-pipeline accuracy ranged from 78.57% to 84.10%. The 979 images are distributed across outer test folds, unlike U-Net’s repeated fixed-test evaluations. Fold-level per-class pipeline metrics and per-image predictions are not contained in the saved DeepLabV3 CV summary CSV.')
p('### 5.5 U-Net pipeline stage-wise fold summary')
table(['Stage','Precision (%)','Recall (%)','F1 (%)'],[[c.title()]+[col(up,c+'_'+k,100,2) for k in ['precision','recall','f1']] for c in classes])
p('These class-wise values are available for U-Net’s fixed-classifier evaluations only. They should not be paired with DeepLabV3’s single-run class metrics as if both were cross-validated class results.')
p('## 6. Interpretation and limits')
p('Segmentation and classification measure different outcomes. U-Net’s stronger foreground overlap in the saved comparison did not automatically yield stronger classification with native preprocessing: the native pipeline scored 78.75%, whereas the crop-aligned variant scored 89.38%. Thus, the saved experiments support evaluating input framing and mask/classifier compatibility in addition to overlap metrics.')
p('The fixed-split pipeline records show 89.38% accuracy for the main U-Net run and 87.50% for the later DeepLabV3 recomputation. Earlier DeepLabV3 comparison results show 84.38%. These are distinct saved evaluations; selecting one score without its protocol would conceal run differences. The report preserves them rather than treating them as duplicates.')
p('The CV studies answer different questions. U-Net quantifies segmenter and downstream variation with a fixed classifier. DeepLabV3 full-pipeline CV quantifies the complete training procedure with both models retrained. Their numerical difference cannot establish a cross-validated architecture advantage because the splits, test exposure, training schedules, and retraining scope differ.')
p('Image-level splits have not been established here as independent by plant, collection session, or near-duplicate group. The original test set was pooled into DeepLabV3 CV and is therefore not an untouched external test for that study. Repeated U-Net fixed-test scoring must not be used for test-guided model selection while claiming untouched evaluation. Cross-project data separation, preprocessing implementation, and duplicate handling require a dedicated audit before stronger generalization claims.')
p('Some saved evaluations differ in resize/crop geometry, soft versus binary Dice, foreground versus both-class metrics, and per-image versus pooled-pixel or batch aggregation. The earlier U-Net report also identifies a mask-resizing concern in the pipeline notebook. These differences should be resolved by a common evaluation implementation before a definitive segmentation comparison. No significance test or external validation is claimed here.')
p('## 7. Conclusions and next reporting steps')
p('Both U-Net and DeepLabV3 have saved segmentation IoU and Dice results in fixed-split evaluations. Both sequential pipelines have saved growth-stage classification results. U-Net has five-fold segmenter training with fixed-classifier downstream evaluation; DeepLabV3 has both segmenter-only CV and full-pipeline CV with both networks retrained.')
p('For presentation, use Section 2 for standalone segmentation, Section 3 for fixed-split pipelines, Section 4 for segmentation CV, and Section 5 for pipeline CV. Keep the protocol note below the pipeline CV table. Do not mix a fixed-split IoU with a cross-validated pipeline score in a row implying one experiment.')
p('To obtain a matched comparison, use identical eligible images and outer folds, harmonize image geometry and overlap definitions, retrain both pipeline models within each fold, select checkpoints only on inner validation, and retain per-image predictions and pixel confusion counts. This is future work, not a completed experiment. DeepLabV3 foreground Dice per CV fold requires additional saved confusion data or reevaluation; it is not recoverable from the current macro summary alone.')
p('## 8. Evidence index and reproducibility')
sources=[
('S1','Original U-Net metadata',ROOT/'export/metadata.json'),
('S2','Fixed-split cross-model comparison',OTHER/'results/unet_vs_deeplabv3_comparison_v2.csv'),
('S3','U-Net fixed-split pipeline predictions',ROOT/'results/fusion_pipeline/full_pipeline_test_results.csv'),
('S4','U-Net segmentation fold metrics',ROOT/'results/cross_validation/unet_segmentation_cv_results.csv'),
('S5','DeepLabV3 segmenter-only CV',OTHER/'results/deeplabv3_kfold_cv_results.csv'),
('S6','DeepLabV3 full-pipeline CV',OTHER/'results/deeplabv3_pipeline_kfold_cv_results.csv'),
('S7','DeepLabV3 training and CV notebook',OTHER/'notebooks/03_deeplabv3_segmentation_pipeline.ipynb'),
('S8','DeepLabV3 standalone pipeline recomputation',OTHER/'notebooks/05_deeplabv3_pipeline_evaluation.ipynb'),
('S9','U-Net CV notebook',ROOT/'notebooks/03_unet_cross_validation.ipynb'),
('S10','Earlier U-Net methods/results audit',ROOT/'reports/research_workload_report.md'),
('S11','U-Net pipeline fold metrics',ROOT/'results/cross_validation/fusion_pipeline_cv_results.csv')]
for sid,label,path in sources: p(f'{sid}. {label}: {path.as_posix()}')
p('Notebook cell numbers in this report are zero-based indices in the saved notebook JSON. Section 5 U-Net numerical tables use S11. The companion build_results_report.py regenerates the CSV-based calculations and both document formats. Some fixed-split values are explicitly transcribed from the identified saved notebook outputs. No external literature claims are required for this results-only report.')

md='\n'.join(parts)
stem=OUT/'segmentation_and_pipeline_results_report'
stem.with_suffix('.md').write_text(md,encoding='utf-8')
doc=Document()
sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.65)
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
normal=doc.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(10)
normal.paragraph_format.space_after=Pt(6)
for name in ['Title','Heading 1','Heading 2']:
    doc.styles[name].font.color.rgb=RGBColor.from_string('214E78')
lines=md.splitlines(); i=0
while i<len(lines):
    line=lines[i]
    if line.startswith('| '):
        rows=[]
        while i<len(lines) and lines[i].startswith('| '):
            vals=[x.strip() for x in lines[i].strip('|').split('|')]
            if not all(x=='---' for x in vals): rows.append(vals)
            i+=1
        t=doc.add_table(rows=1,cols=len(rows[0])); t.style='Light Shading Accent 1'
        for j,val in enumerate(rows[0]): t.rows[0].cells[j].text=val
        repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows[1:]:
            cells=t.add_row().cells
            for j,val in enumerate(row): cells[j].text=val
        for row in t.rows:
            prop=row._tr.get_or_add_trPr(); prop.append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.space_after=Pt(3)
                    for run in para.runs: run.font.size=Pt(8)
        doc.add_paragraph()
        continue
    elif line.startswith('# '): doc.add_heading(line[2:],0)
    elif line.startswith('## '): doc.add_heading(line[3:],1)
    elif line.startswith('### '): doc.add_heading(line[4:],2)
    elif line: doc.add_paragraph(line)
    i+=1
footer=sec.footer.paragraphs[0]; footer.alignment=2
footer.add_run('Results report | ')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
doc.save(stem.with_suffix('.docx'))
print('Created',stem.with_suffix('.md'))
print('Created',stem.with_suffix('.docx'))
print('Tables',len(doc.tables),'Words',len(md.split()))
assert len(u)==len(up)==len(d)==len(ds)==5
assert len(preds)==160 and sum(cm[i][i] for i in range(4))==143
check=Document(stem.with_suffix('.docx'))
assert len(check.tables)==len(doc.tables)
assert '84.00 ± 2.39' in md and '77.12 ± 6.74' in md
print('Verified fold counts, prediction counts, key summaries, and DOCX readability.')
